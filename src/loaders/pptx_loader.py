from typing import List

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.loaders.base import LoadContext, LoaderError, image_element, is_ole_container, register, rows_to_markdown
from src.schema import Element


def iter_shapes(shapes):
    """Yields shapes in order, descending into group shapes."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)
        else:
            yield shape


def shape_image_blob(shape):
    """Returns picture bytes for pictures and picture placeholders, else None."""
    try:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
            return shape.image.blob
    except Exception:
        return None
    return None


def chart_rows(chart) -> List[List]:
    """Categories and series values as a table: header = ["Series", *categories]."""
    categories = []
    try:
        categories = [str(c) for c in chart.plots[0].categories] if len(chart.plots) else []
    except Exception:
        categories = []
    rows = [["Series"] + categories]
    for plot in chart.plots:
        for series in plot.series:
            try:
                values = list(series.values)
            except Exception:
                values = []
            rows.append([series.name or ""] + values)
    return rows


def chart_title(chart) -> str:
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            return chart.chart_title.text_frame.text.strip()
    except Exception:
        pass
    return ""


@register(".pptx")
def load_pptx(ctx: LoadContext) -> List[Element]:
    """Per slide: text, tables, charts, pictures (vision) and speaker notes, each citing its slide number."""
    if is_ole_container(ctx.path):
        raise LoaderError("file is password protected")
    try:
        presentation = Presentation(ctx.path)
    except Exception as e:
        raise LoaderError(f"file is corrupt or not a PowerPoint file ({type(e).__name__})")

    elements: List[Element] = []
    pending = []  # ((image_bytes, context, kind), element kwargs)
    empty_slides = []

    for slide_no, slide in enumerate(presentation.slides, start=1):
        loc = {"slide": slide_no}
        label = f"slide {slide_no}"
        hidden = slide._element.get("show") == "0"
        extra = {"hidden": True} if hidden else {}
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        heading_path = [title] if title else []

        texts, slide_elements, images = [], [], []
        for shape in iter_shapes(slide.shapes):
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    texts.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                rows = [r for r in rows if any(r)]
                if rows:
                    slide_elements.append(
                        ctx.make(
                            "table",
                            rows_to_markdown(rows[0], rows[1:]),
                            loc,
                            label,
                            heading_path=heading_path,
                            title=title,
                            structured={"header": rows[0], "rows": rows[1:]},
                            extra=extra,
                        )
                    )
            if getattr(shape, "has_chart", False) and shape.has_chart:
                rows = chart_rows(shape.chart)
                name = chart_title(shape.chart)
                body = rows_to_markdown(rows[0], rows[1:])
                slide_elements.append(
                    ctx.make(
                        "chart",
                        (f"Chart: {name}\n\n" if name else "") + body,
                        loc,
                        label,
                        heading_path=heading_path,
                        title=title,
                        structured={"chart_title": name, "header": rows[0], "rows": rows[1:]},
                        extra=extra,
                    )
                )
            blob = shape_image_blob(shape)
            if blob:
                images.append(blob)

        slide_text = "\n".join(texts)
        if slide_text:
            elements.append(ctx.make("text", slide_text, loc, label, heading_path=heading_path, title=title, extra=extra))
        elements.extend(slide_elements)

        for blob in images:
            if ctx.take_image_slot():
                pending.append(
                    (
                        (blob, slide_text[:1500], "image"),
                        dict(location=loc, label=label, heading_path=heading_path, title=title, context_text=slide_text, extra=extra),
                    )
                )

        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            elements.append(
                ctx.make(
                    "text",
                    notes,
                    loc,
                    f"slide {slide_no} speaker notes",
                    heading_path=heading_path,
                    title=title,
                    extra={**extra, "notes": True},
                )
            )

        if not slide_text and not slide_elements and not images and not notes:
            empty_slides.append(slide_no)

    for (_, kwargs), (description, error) in zip(pending, ctx.describe_images([job for job, _ in pending])):
        elements.append(image_element(ctx, description, error, **kwargs))

    if empty_slides:
        ctx.warnings.append(f"no readable content on slide(s) {', '.join(map(str, empty_slides))}")
    elements.sort(key=lambda e: (e.metadata["slide"], e.element_id))
    return elements
