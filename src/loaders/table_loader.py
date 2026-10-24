import csv
from typing import List, Optional, Tuple

import pandas as pd

from src.loaders.base import LoadContext, LoaderError, is_ole_container, register, rows_to_markdown
from src.schema import Element

# A sheet is a list of (sheet_row_number, values) with 1-based row numbers as the user sees them.
Rows = List[Tuple[int, List]]


def _is_empty(values) -> bool:
    return all(v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == "" for v in values)


def sheet_elements(ctx: LoadContext, sheet: Optional[str], rows: Rows) -> List[Element]:
    """First non-empty row is the header; data rows are batched into Markdown tables that repeat it."""
    rows = [(n, v) for n, v in rows if not _is_empty(v)]
    if not rows:
        return []
    header_no, header = rows[0]
    data = rows[1:]
    cap = ctx.settings.max_rows_per_sheet
    if len(data) > cap:
        ctx.warnings.append(f"{len(data) - cap} row(s) skipped in {sheet or 'file'} (limit {cap} rows per sheet)")
        data = data[:cap]
    if not data:
        data = []

    elements = []
    size = ctx.settings.rows_per_element
    batches = [data[i : i + size] for i in range(0, len(data), size)] or [[]]
    for batch in batches:
        row_start = batch[0][0] if batch else header_no
        row_end = batch[-1][0] if batch else header_no
        values = [v for _, v in batch]
        location = {"row_start": row_start, "row_end": row_end}
        label = f"rows {row_start}-{row_end}"
        if sheet is not None:
            location["sheet"] = sheet
            label = f"{sheet}, {label}"
        elements.append(
            ctx.make(
                "table",
                rows_to_markdown(header, values),
                location,
                label,
                title=sheet or "",
                heading_path=[sheet] if sheet else [],
                structured={"header": [str(h) for h in header], "rows": [[_plain(c) for c in r] for r in values]},
            )
        )
    return elements


def _plain(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (int, float, str, bool)):
        return value
    return str(value)


@register(".csv")
def load_csv(ctx: LoadContext) -> List[Element]:
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            with open(ctx.path, newline="", encoding=encoding) as f:
                sample = f.read(64 * 1024)
                f.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                rows = [(n, values) for n, values in enumerate(csv.reader(f, dialect), start=1)]
            break
        except UnicodeDecodeError:
            continue
    else:
        raise LoaderError("could not decode CSV")
    return sheet_elements(ctx, None, rows)


@register(".xlsx")
def load_xlsx(ctx: LoadContext) -> List[Element]:
    """Reads every sheet with openpyxl; merged cells repeat their value in every covered cell."""
    from openpyxl import load_workbook

    if is_ole_container(ctx.path):
        raise LoaderError("file is password protected")
    try:
        # read_only mode has no merged-cell info, so load normally (file size is capped upstream)
        workbook = load_workbook(ctx.path, data_only=True)
    except Exception as e:
        raise LoaderError(f"file is corrupt or not an Excel workbook ({type(e).__name__})")

    elements = []
    for ws in workbook.worksheets:
        grid = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    grid[(cell.row, cell.column)] = cell.value
        for merged in ws.merged_cells.ranges:
            value = grid.get((merged.min_row, merged.min_col))
            for r in range(merged.min_row, merged.max_row + 1):
                for c in range(merged.min_col, merged.max_col + 1):
                    grid[(r, c)] = value
        if not grid:
            continue
        max_row = max(r for r, _ in grid)
        max_col = max(c for _, c in grid)
        rows = [(r, [grid.get((r, c)) for c in range(1, max_col + 1)]) for r in range(1, max_row + 1)]
        elements.extend(sheet_elements(ctx, ws.title, rows))
    workbook.close()
    return elements


@register(".xls")
def load_xls(ctx: LoadContext) -> List[Element]:
    try:
        sheets = pd.read_excel(ctx.path, sheet_name=None, header=None, engine="xlrd")
    except Exception as e:
        if "encrypt" in str(e).lower() or "password" in str(e).lower():
            raise LoaderError("file is password protected")
        raise LoaderError(f"file is corrupt or not an Excel workbook ({type(e).__name__})")
    elements = []
    for name, df in sheets.items():
        rows = [(i + 1, list(values)) for i, values in enumerate(df.itertuples(index=False, name=None))]
        elements.extend(sheet_elements(ctx, str(name), rows))
    return elements
