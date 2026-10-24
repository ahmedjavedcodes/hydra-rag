import io
import os

import pytest
from PIL import Image, ImageDraw

import utilities.session as session_module
from src.loaders.vision import VisionDescriber

FAKE_DESCRIPTION = "This image contains a blue rectangle used as a test figure."


@pytest.fixture(autouse=True)
def data_root(tmp_path, monkeypatch):
    """Every test writes to its own ./data replacement."""
    root = tmp_path / "data"
    monkeypatch.setattr(session_module, "DATA_ROOT", str(root))
    return root


@pytest.fixture
def vision():
    """Vision describer that never calls a real model."""
    return VisionDescriber(invoke=lambda prompt, image_url: FAKE_DESCRIPTION, wait_min=0, wait_max=0)


def png_bytes(width=200, height=150, color=(30, 60, 200)) -> bytes:
    img = Image.new("RGB", (width, height), color)
    ImageDraw.Draw(img).text((10, 10), "Q3 revenue", fill=(255, 255, 255))
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.fixture
def samples(tmp_path):
    """Builds one sample file per supported format and returns their paths by key."""
    import pymupdf as fitz
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    out = tmp_path / "inputs"
    out.mkdir()
    paths = {}

    # PDF: two pages, an outline, a ruled table and an embedded image
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Introduction to the quarterly report.", fontsize=12)
    page.insert_text((72, 90), "Revenue grew strongly in every region this quarter.", fontsize=12)
    x0, y0, cw, rh = 72, 140, 120, 24
    cells = [["Region", "Revenue", "Growth"], ["North", "120", "5%"], ["South", "95", "3%"]]
    for r, row in enumerate(cells):
        for c, value in enumerate(row):
            page.insert_text((x0 + c * cw + 6, y0 + r * rh + 16), value, fontsize=10)
    for r in range(len(cells) + 1):
        page.draw_line((x0, y0 + r * rh), (x0 + 3 * cw, y0 + r * rh))
    for c in range(4):
        page.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + len(cells) * rh))
    page.insert_text((72, 300), "Figure 1: revenue by region.", fontsize=10)
    page.insert_image(fitz.Rect(72, 310, 272, 460), stream=png_bytes())
    page2 = pdf.new_page()
    page2.insert_text((72, 72), "Results: the North region led growth at five percent.", fontsize=12)
    pdf.set_toc([[1, "Introduction", 1], [1, "Results", 2]])
    paths["pdf"] = str(out / "report.pdf")
    pdf.save(paths["pdf"])
    pdf.close()

    # Scanned PDF: one page that is only an image (no text layer)
    scanned = fitz.open()
    scanned.new_page().insert_image(fitz.Rect(0, 0, 595, 842), stream=png_bytes(600, 800, (240, 240, 240)))
    paths["scanned_pdf"] = str(out / "scanned.pdf")
    scanned.save(paths["scanned_pdf"])
    scanned.close()

    # Encrypted PDF
    locked = fitz.open()
    locked.new_page().insert_text((72, 72), "secret", fontsize=12)
    paths["locked_pdf"] = str(out / "locked.pdf")
    locked.save(paths["locked_pdf"], encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="o", user_pw="u")
    locked.close()

    # DOCX: headings, paragraphs, a table
    doc = Document()
    doc.add_heading("Overview", level=1)
    doc.add_paragraph("HydraRAG answers questions about uploaded documents.")
    doc.add_heading("Pricing", level=2)
    doc.add_paragraph("The basic plan costs ten dollars per month.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Plan", "Price"
    table.cell(1, 0).text, table.cell(1, 1).text = "Basic", "$10"
    paths["docx"] = str(out / "guide.docx")
    doc.save(paths["docx"])

    # PPTX: title, bullets, table, chart, picture, notes, hidden slide
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "HydraRAG Launch"
    s1.placeholders[1].text = "Quarterly update"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Goals"
    s2.placeholders[1].text = "Ship ingestion\nAdd Pinecone"
    s2.notes_slide.notes_text_frame.text = "Mention the WhatsApp channel here."
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Budget"
    shape = s3.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1))
    shape.table.cell(0, 0).text, shape.table.cell(0, 1).text = "Item", "Cost"
    shape.table.cell(1, 0).text, shape.table.cell(1, 1).text = "Servers", "300"
    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    s4.shapes.title.text = "Growth"
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2"]
    chart_data.add_series("Users", (100, 250))
    s4.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), chart_data)
    s5 = prs.slides.add_slide(prs.slide_layouts[5])
    s5.shapes.title.text = "Architecture"
    s5.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(2))
    s6 = prs.slides.add_slide(prs.slide_layouts[5])
    s6.shapes.title.text = "Backup slide"
    s6._element.set("show", "0")
    paths["pptx"] = str(out / "deck.pptx")
    prs.save(paths["pptx"])

    # XLSX: two sheets, a merged cell
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Q1", "Q2"])
    ws.append(["North", 10, 12])
    ws.append([None, None, None])
    ws.append(["South", 7, 9])
    ws2 = wb.create_sheet("Staff")
    ws2.append(["Team", "Name"])
    ws2.append(["Core", "Ana"])
    ws2.append([None, "Bo"])
    ws2.merge_cells("A2:A3")
    paths["xlsx"] = str(out / "numbers.xlsx")
    wb.save(paths["xlsx"])

    paths["csv"] = str(out / "data.csv")
    with open(paths["csv"], "w", encoding="utf-8", newline="") as f:
        f.write("city,population\nLahore,13000000\nKarachi,16000000\n")

    paths["txt"] = str(out / "notes.txt")
    with open(paths["txt"], "w", encoding="utf-8") as f:
        f.write("First paragraph about ingestion.\n\nSecond paragraph about retrieval.\n")

    paths["md"] = str(out / "readme.md")
    with open(paths["md"], "w", encoding="utf-8") as f:
        f.write("# Title\n\nIntro text.\n\n## Setup\n\nRun uv sync.\n")

    paths["png"] = str(out / "chart.png")
    with open(paths["png"], "wb") as f:
        f.write(png_bytes())

    paths["corrupt_pdf"] = str(out / "broken.pdf")
    with open(paths["corrupt_pdf"], "wb") as f:
        f.write(b"%PDF-1.4\nthis is not really a pdf")

    paths["exe"] = str(out / "tool.exe")
    with open(paths["exe"], "wb") as f:
        f.write(b"MZ\x90\x00")

    paths["ppt"] = str(out / "old.ppt")
    with open(paths["ppt"], "wb") as f:
        f.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)

    paths["locked_docx"] = str(out / "locked.docx")
    with open(paths["locked_docx"], "wb") as f:
        f.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)

    paths["empty_txt"] = str(out / "empty.txt")
    open(paths["empty_txt"], "w").close()

    assert all(os.path.exists(p) for p in paths.values())
    return paths
